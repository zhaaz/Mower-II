"""Integriertes Testfenster für den Murata SCH16T-K01 Datenstrom."""

from __future__ import annotations

import math
import queue
import threading
from dataclasses import dataclass
from datetime import datetime
from pathlib import Path
import tkinter as tk
from tkinter import filedialog, messagebox, ttk

import serial
from serial.tools import list_ports


MICROS_WRAP = 2**32


@dataclass(frozen=True)
class Measurement:
    pc_time: datetime
    sensor_time_s: float
    gyro_z_dps: float
    yaw_raw_deg: float
    yaw_relative_deg: float
    drift_deg: float
    yaw_corrected_deg: float


class _SerialReader(threading.Thread):
    def __init__(self, port: str, baudrate: int, output_queue: queue.Queue) -> None:
        super().__init__(daemon=True)
        self.port = port
        self.baudrate = int(baudrate)
        self.output_queue = output_queue
        self.stop_event = threading.Event()
        self.connection: serial.Serial | None = None
        self.last_time_us: int | None = None
        self.wrap_offset_us = 0

    def stop(self) -> None:
        self.stop_event.set()
        connection = self.connection
        if connection is not None and connection.is_open:
            try:
                connection.close()
            except (OSError, AttributeError, serial.SerialException):
                pass

    def _extend_micros(self, current: int) -> int:
        if (
            self.last_time_us is not None
            and current < self.last_time_us
            and self.last_time_us - current > MICROS_WRAP // 2
        ):
            self.wrap_offset_us += MICROS_WRAP
        self.last_time_us = current
        return current + self.wrap_offset_us

    def run(self) -> None:
        try:
            # DTR/RTS bereits vor dem Öffnen deaktivieren. Dadurch loest das
            # Diagnosefenster beim Nano ESP32 keinen unnoetigen Reset aus.
            connection = serial.Serial()
            connection.port = self.port
            connection.baudrate = self.baudrate
            connection.timeout = 0.25
            connection.dtr = False
            connection.rts = False
            connection.open()
            self.connection = connection
            connection.reset_input_buffer()
            self.output_queue.put(("connected", self.port))

            while not self.stop_event.is_set():
                raw_line = connection.readline()
                if not raw_line:
                    continue

                line = raw_line.decode("utf-8", errors="replace").strip()
                if not line or line.startswith("#") or line.startswith("time_us"):
                    continue

                fields = line.split(";")
                if len(fields) != 3:
                    self.output_queue.put(("invalid",))
                    continue

                try:
                    raw_time_us = int(fields[0])
                    gyro_z_dps = float(fields[1].replace(",", "."))
                    yaw_raw_deg = float(fields[2].replace(",", "."))
                except ValueError:
                    self.output_queue.put(("invalid",))
                    continue

                self.output_queue.put(
                    (
                        "sample",
                        datetime.now(),
                        self._extend_micros(raw_time_us) / 1_000_000.0,
                        gyro_z_dps,
                        yaw_raw_deg,
                    )
                )
        except (serial.SerialException, OSError) as exc:
            if not self.stop_event.is_set():
                self.output_queue.put(("error", str(exc)))
        finally:
            connection = self.connection
            if connection is not None and connection.is_open:
                try:
                    connection.close()
                except (OSError, AttributeError, serial.SerialException):
                    pass
            self.output_queue.put(("disconnected", self.port))


class MurataTestDialog:
    def __init__(
        self,
        parent: tk.Misc,
        *,
        default_port: str,
        baudrate: int,
    ) -> None:
        self.parent = parent
        self.baudrate = int(baudrate)
        self.window = tk.Toplevel(parent)
        self.window.title("Murata SCH16T Test")
        self.window.geometry("820x620")
        self.window.minsize(760, 560)
        self.window.transient(parent)
        self.window.grab_set()

        self.message_queue: queue.Queue = queue.Queue()
        self.reader: _SerialReader | None = None
        self.closed = False
        self.valid_packets = 0
        self.invalid_lines = 0

        self.latest_sensor_time_s: float | None = None
        self.latest_raw_yaw_deg: float | None = None
        self.latest_corrected_yaw_deg = 0.0
        self.zero_sensor_time_s: float | None = None
        self.zero_raw_yaw_deg = 0.0
        self.drift_rate_dps = 0.0
        self.drift_duration_actual_s = 0.0
        self.drift_active = False
        self.drift_duration_requested_s = 10.0
        self.drift_start_sensor_time_s: float | None = None
        self.drift_points: list[tuple[float, float]] = []
        self.recording = False
        self.measurements: list[Measurement] = []

        self.port_var = tk.StringVar(value=default_port)
        self.duration_var = tk.StringVar(value="10")
        self.status_var = tk.StringVar(value="Nicht verbunden")
        self.raw_var = tk.StringVar(value="–")
        self.corrected_var = tk.StringVar(value="–")
        self.gyro_var = tk.StringVar(value="–")
        self.drift_var = tk.StringVar(value="Noch nicht gemessen")
        self.record_var = tk.StringVar(value="0 Datensätze")
        self.packet_var = tk.StringVar(value="Gültig: 0 | Ungültig: 0")

        self._build_ui()
        self.refresh_ports()
        self.window.protocol("WM_DELETE_WINDOW", self.close)
        self.window.bind("<Escape>", lambda _event: self.close())
        self.window.after(20, self.process_queue)

    def _build_ui(self) -> None:
        main = ttk.Frame(self.window, padding=14)
        main.pack(fill="both", expand=True)
        main.columnconfigure(0, weight=1)
        main.columnconfigure(1, weight=1)
        main.rowconfigure(3, weight=1)

        connection = ttk.LabelFrame(main, text="Verbindung", padding=10)
        connection.grid(row=0, column=0, columnspan=2, sticky="ew", pady=(0, 10))
        connection.columnconfigure(1, weight=1)
        ttk.Label(connection, text="COM-Port:").grid(row=0, column=0, padx=(0, 8))
        self.port_box = ttk.Combobox(
            connection, textvariable=self.port_var, state="readonly", width=26
        )
        self.port_box.grid(row=0, column=1, sticky="ew")
        ttk.Button(connection, text="Aktualisieren", command=self.refresh_ports).grid(
            row=0, column=2, padx=8
        )
        self.connect_button = ttk.Button(
            connection, text="Verbinden", command=self.toggle_connection
        )
        self.connect_button.grid(row=0, column=3)
        ttk.Label(connection, text=f"Baudrate: {self.baudrate}").grid(
            row=1, column=0, sticky="w", pady=(8, 0)
        )
        ttk.Label(connection, textvariable=self.status_var).grid(
            row=1, column=1, columnspan=2, sticky="w", pady=(8, 0)
        )
        ttk.Label(connection, textvariable=self.packet_var).grid(
            row=1, column=3, sticky="e", pady=(8, 0)
        )

        controls = ttk.LabelFrame(main, text="Messung", padding=10)
        controls.grid(row=1, column=0, columnspan=2, sticky="ew", pady=(0, 10))
        ttk.Label(controls, text="Driftdauer [s]:").grid(row=0, column=0, padx=(0, 6))
        ttk.Entry(controls, textvariable=self.duration_var, width=8).grid(
            row=0, column=1, padx=(0, 8)
        )
        self.drift_button = ttk.Button(
            controls, text="Drift messen", command=self.start_drift_measurement
        )
        self.drift_button.grid(row=0, column=2, padx=(0, 8))
        ttk.Button(controls, text="Richtung = 0°", command=self.set_zero).grid(
            row=0, column=3, padx=(0, 8)
        )
        self.record_button = ttk.Button(
            controls, text="Aufzeichnung starten", command=self.toggle_recording
        )
        self.record_button.grid(row=0, column=4, padx=(0, 8))
        ttk.Button(controls, text="Excel speichern", command=self.save_excel).grid(
            row=0, column=5
        )

        values = ttk.LabelFrame(main, text="Aktuelle Werte", padding=10)
        values.grid(row=2, column=0, columnspan=2, sticky="ew", pady=(0, 10))
        for column in range(4):
            values.columnconfigure(column, weight=1)
        self._value_field(values, 0, "Gyro Z", self.gyro_var)
        self._value_field(values, 1, "Yaw roh", self.raw_var)
        self._value_field(values, 2, "Yaw korrigiert", self.corrected_var)
        self._value_field(values, 3, "Aufzeichnung", self.record_var)

        direction = ttk.LabelFrame(main, text="Richtung (korrigiert)", padding=8)
        direction.grid(row=3, column=0, sticky="nsew", padx=(0, 5))
        direction.rowconfigure(0, weight=1)
        direction.columnconfigure(0, weight=1)
        self.canvas = tk.Canvas(direction, background="white", highlightthickness=0)
        self.canvas.grid(row=0, column=0, sticky="nsew")
        self.canvas.bind(
            "<Configure>",
            lambda _event: self.draw_direction(self.latest_corrected_yaw_deg),
        )

        drift = ttk.LabelFrame(main, text="Drift", padding=12)
        drift.grid(row=3, column=1, sticky="nsew", padx=(5, 0))
        drift.columnconfigure(0, weight=1)
        ttk.Label(
            drift, textvariable=self.drift_var, justify="left", font=("Segoe UI", 11)
        ).grid(row=0, column=0, sticky="nw")
        self.progress = ttk.Progressbar(drift, maximum=100, mode="determinate")
        self.progress.grid(row=1, column=0, sticky="ew", pady=(16, 8))
        ttk.Label(
            drift,
            text=(
                "Während der Driftmessung muss der Sensor vollständig ruhig stehen.\n"
                "Die Korrektur wird ab dem gesetzten Nullpunkt angewendet."
            ),
            justify="left",
            wraplength=320,
        ).grid(row=2, column=0, sticky="nw")
        ttk.Button(drift, text="Schließen", command=self.close).grid(
            row=3, column=0, sticky="se", pady=(24, 0)
        )

    @staticmethod
    def _value_field(parent: tk.Misc, column: int, title: str, variable: tk.StringVar) -> None:
        frame = ttk.Frame(parent, padding=(6, 2))
        frame.grid(row=0, column=column, sticky="ew")
        ttk.Label(frame, text=title).pack(anchor="w")
        ttk.Label(frame, textvariable=variable, font=("Segoe UI", 14, "bold")).pack(anchor="w")

    def refresh_ports(self) -> None:
        ports = [item.device for item in list_ports.comports()]
        current = self.port_var.get().strip()
        if current and current not in ports:
            ports.insert(0, current)
        self.port_box.configure(values=ports)
        if not current and ports:
            self.port_var.set(ports[0])

    def toggle_connection(self) -> None:
        if self.reader is not None and self.reader.is_alive():
            self.reader.stop()
            self.status_var.set("Verbindung wird getrennt ...")
            return
        port = self.port_var.get().strip()
        if not port:
            messagebox.showwarning("Kein COM-Port", "Bitte einen COM-Port auswählen.", parent=self.window)
            return
        self.reader = _SerialReader(port, self.baudrate, self.message_queue)
        self.reader.start()
        self.status_var.set(f"Verbinde mit {port} ...")
        self.connect_button.configure(text="Trennen")
        self.port_box.configure(state="disabled")

    def process_queue(self) -> None:
        if self.closed:
            return
        try:
            while True:
                message = self.message_queue.get_nowait()
                kind = message[0]
                if kind == "connected":
                    self.status_var.set(f"Verbunden: {message[1]} @ {self.baudrate} Baud")
                elif kind == "disconnected":
                    self.status_var.set("Nicht verbunden")
                    self.connect_button.configure(text="Verbinden")
                    self.port_box.configure(state="readonly")
                    self.recording = False
                    self.record_button.configure(text="Aufzeichnung starten")
                elif kind == "error":
                    messagebox.showerror("Serieller Fehler", message[1], parent=self.window)
                elif kind == "invalid":
                    self.invalid_lines += 1
                    self._update_packet_text()
                elif kind == "sample":
                    self.handle_sample(*message[1:])
        except queue.Empty:
            pass
        self.window.after(20, self.process_queue)

    def _update_packet_text(self) -> None:
        self.packet_var.set(f"Gültig: {self.valid_packets} | Ungültig: {self.invalid_lines}")

    def handle_sample(
        self,
        pc_time: datetime,
        sensor_time_s: float,
        gyro_z_dps: float,
        yaw_raw_deg: float,
    ) -> None:
        self.valid_packets += 1
        self._update_packet_text()
        self.latest_sensor_time_s = sensor_time_s
        self.latest_raw_yaw_deg = yaw_raw_deg
        if self.zero_sensor_time_s is None:
            self.zero_sensor_time_s = sensor_time_s
            self.zero_raw_yaw_deg = yaw_raw_deg

        elapsed = sensor_time_s - self.zero_sensor_time_s
        yaw_relative = yaw_raw_deg - self.zero_raw_yaw_deg
        drift_deg = self.drift_rate_dps * elapsed
        yaw_corrected = yaw_relative - drift_deg
        self.latest_corrected_yaw_deg = yaw_corrected
        self.gyro_var.set(f"{gyro_z_dps:+.5f} °/s")
        self.raw_var.set(f"{yaw_raw_deg:+.4f}°")
        self.corrected_var.set(f"{yaw_corrected:+.4f}°")
        self.draw_direction(yaw_corrected)

        if self.drift_active:
            self.update_drift_measurement(sensor_time_s, yaw_raw_deg)
        if self.recording:
            self.measurements.append(
                Measurement(
                    pc_time, sensor_time_s, gyro_z_dps, yaw_raw_deg,
                    yaw_relative, drift_deg, yaw_corrected,
                )
            )
            self.record_var.set(f"{len(self.measurements)} Datensätze")

    def start_drift_measurement(self) -> None:
        if self.latest_sensor_time_s is None:
            messagebox.showwarning("Keine Daten", "Zuerst verbinden und Daten abwarten.", parent=self.window)
            return
        try:
            duration = float(self.duration_var.get().replace(",", "."))
            if duration <= 0.0:
                raise ValueError
        except ValueError:
            messagebox.showwarning("Ungültige Dauer", "Bitte eine positive Dauer eingeben.", parent=self.window)
            return
        self.drift_duration_requested_s = duration
        self.drift_start_sensor_time_s = None
        self.drift_points.clear()
        self.drift_active = True
        self.progress.configure(value=0)
        self.drift_button.configure(state="disabled")

    def update_drift_measurement(self, sensor_time_s: float, yaw_raw_deg: float) -> None:
        if self.drift_start_sensor_time_s is None:
            self.drift_start_sensor_time_s = sensor_time_s
        elapsed = sensor_time_s - self.drift_start_sensor_time_s
        self.drift_points.append((elapsed, yaw_raw_deg))
        self.progress.configure(value=min(100.0, elapsed / self.drift_duration_requested_s * 100.0))
        self.drift_var.set(f"Messung läuft: {elapsed:.1f} / {self.drift_duration_requested_s:.1f} s")
        if elapsed >= self.drift_duration_requested_s:
            self.finish_drift_measurement()

    def finish_drift_measurement(self) -> None:
        self.drift_active = False
        self.drift_button.configure(state="normal")
        self.progress.configure(value=100)
        if len(self.drift_points) < 2:
            self.drift_var.set("Driftmessung fehlgeschlagen: zu wenige Werte")
            return
        times = [item[0] for item in self.drift_points]
        angles = [item[1] for item in self.drift_points]
        mean_time = sum(times) / len(times)
        mean_angle = sum(angles) / len(angles)
        denominator = sum((value - mean_time) ** 2 for value in times)
        if denominator == 0.0:
            self.drift_var.set("Driftmessung fehlgeschlagen: keine Zeitspanne")
            return
        self.drift_rate_dps = sum(
            (sample_time - mean_time) * (angle - mean_angle)
            for sample_time, angle in self.drift_points
        ) / denominator
        self.drift_duration_actual_s = times[-1] - times[0]
        self.drift_var.set(
            "Driftmessung abgeschlossen\n"
            f"{self.drift_rate_dps:+.7f} °/s\n"
            f"{self.drift_rate_dps * 60:+.5f} °/min\n"
            f"{self.drift_rate_dps * 3600:+.3f} °/h\n"
            f"{len(self.drift_points)} Messwerte"
        )

    def set_zero(self) -> None:
        if self.latest_sensor_time_s is None or self.latest_raw_yaw_deg is None:
            messagebox.showwarning("Keine Daten", "Es liegt noch kein Sensorwert vor.", parent=self.window)
            return
        self.zero_sensor_time_s = self.latest_sensor_time_s
        self.zero_raw_yaw_deg = self.latest_raw_yaw_deg
        self.latest_corrected_yaw_deg = 0.0
        self.corrected_var.set("+0.0000°")
        self.draw_direction(0.0)

    def toggle_recording(self) -> None:
        if not self.recording:
            if self.latest_sensor_time_s is None:
                messagebox.showwarning("Keine Daten", "Zuerst verbinden und Daten abwarten.", parent=self.window)
                return
            self.measurements.clear()
            self.recording = True
            self.record_button.configure(text="Aufzeichnung stoppen")
            self.record_var.set("0 Datensätze")
        else:
            self.recording = False
            self.record_button.configure(text="Aufzeichnung starten")
            self.record_var.set(f"{len(self.measurements)} Datensätze")

    def draw_direction(self, yaw_deg: float) -> None:
        self.canvas.delete("all")
        width = max(self.canvas.winfo_width(), 260)
        height = max(self.canvas.winfo_height(), 260)
        center_x, center_y = width / 2, height / 2
        radius = max(70.0, min(width, height) * 0.34)
        self.canvas.create_oval(
            center_x - radius, center_y - radius,
            center_x + radius, center_y + radius,
            width=2, outline="#536273",
        )
        for angle, label in ((0, "0°"), (90, "90°"), (180, "180°"), (270, "270°")):
            radians = math.radians(angle)
            x = center_x + math.sin(radians) * (radius + 20)
            y = center_y - math.cos(radians) * (radius + 20)
            self.canvas.create_text(x, y, text=label, fill="#536273")
        display_angle = yaw_deg % 360.0
        radians = math.radians(display_angle)
        end_x = center_x + math.sin(radians) * (radius - 14)
        end_y = center_y - math.cos(radians) * (radius - 14)
        self.canvas.create_line(
            center_x, center_y, end_x, end_y,
            width=5, fill="#1769aa", arrow=tk.LAST, arrowshape=(18, 22, 8),
        )
        self.canvas.create_text(
            center_x, center_y + radius + 42,
            text=f"{yaw_deg:+.3f}°  ({display_angle:.1f}°)",
            font=("Segoe UI", 13, "bold"), fill="#1f2933",
        )

    def save_excel(self) -> None:
        if not self.measurements:
            messagebox.showwarning("Keine Messdaten", "Zuerst eine Aufzeichnung durchführen.", parent=self.window)
            return
        try:
            from openpyxl import Workbook
            from openpyxl.chart import LineChart, Reference
        except ImportError:
            messagebox.showerror(
                "OpenPyXL fehlt",
                "Installation: python -m pip install openpyxl",
                parent=self.window,
            )
            return

        filename = filedialog.asksaveasfilename(
            parent=self.window,
            title="Messdaten als Excel-Datei speichern",
            defaultextension=".xlsx",
            initialfile=f"SCH16T_Test_{datetime.now():%Y%m%d_%H%M%S}.xlsx",
            filetypes=[("Excel-Arbeitsmappe", "*.xlsx")],
        )
        if not filename:
            return

        workbook = Workbook()
        summary = workbook.active
        summary.title = "Übersicht"
        data_sheet = workbook.create_sheet("Messdaten")
        summary.append(["SCH16T-K01 Yaw-Test"])
        for label, value in (
            ("Exportzeit", datetime.now()),
            ("COM-Port", self.port_var.get()),
            ("Baudrate", self.baudrate),
            ("Messwerte", len(self.measurements)),
            ("Driftrate [°/s]", self.drift_rate_dps),
            ("Driftrate [°/h]", self.drift_rate_dps * 3600.0),
        ):
            summary.append([label, value])

        data_sheet.append([
            "PC-Zeit", "Sensorzeit [s]", "Gyro Z [°/s]", "Yaw roh [°]",
            "Yaw relativ [°]", "Drift [°]", "Yaw korrigiert [°]",
        ])
        for item in self.measurements:
            data_sheet.append([
                item.pc_time, item.sensor_time_s, item.gyro_z_dps, item.yaw_raw_deg,
                item.yaw_relative_deg, item.drift_deg, item.yaw_corrected_deg,
            ])
        chart = LineChart()
        chart.title = "Yaw-Verlauf"
        chart.y_axis.title = "Winkel [°]"
        chart.x_axis.title = "Sensorzeit [s]"
        chart.add_data(
            Reference(data_sheet, min_col=4, max_col=7, min_row=1, max_row=data_sheet.max_row),
            titles_from_data=True,
        )
        chart.set_categories(Reference(data_sheet, min_col=2, min_row=2, max_row=data_sheet.max_row))
        summary.add_chart(chart, "D3")
        try:
            workbook.save(Path(filename))
        except OSError as exc:
            messagebox.showerror("Speichern fehlgeschlagen", str(exc), parent=self.window)
            return
        messagebox.showinfo(
            "Gespeichert",
            f"{len(self.measurements)} Messwerte gespeichert:\n{filename}",
            parent=self.window,
        )

    def close(self) -> None:
        if self.closed:
            return
        self.closed = True
        if self.reader is not None:
            self.reader.stop()
            if self.reader.is_alive():
                self.reader.join(timeout=1.0)
        try:
            self.window.grab_release()
        except tk.TclError:
            pass
        self.window.destroy()


def show_murata_test_dialog(
    parent: tk.Misc,
    *,
    default_port: str,
    baudrate: int,
) -> None:
    dialog = MurataTestDialog(
        parent,
        default_port=default_port,
        baudrate=baudrate,
    )
    parent.wait_window(dialog.window)
